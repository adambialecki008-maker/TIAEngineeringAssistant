"""Read-only XLSX quality reports for HMI batch checkpoints and TIA QA scans.

No import, editing, language activation, project save, or Ollama requests.
Reports rely on confirmed checkpoint outcomes and fresh read-only exports.
"""
from io import BytesIO
from datetime import datetime, timezone

from fastapi import APIRouter
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.tia.hmi_batches import _JOB_LOCK, _load_job, _public_job
from engine.tia.hmi_review import _audit_existing_hmi_translations
from engine.tia.models import TranslateHmiTextsRequest


router = APIRouter()
XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _safe_cell(value):
    """Prevent Excel interpreting untrusted project texts as formulas."""
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


def _sheet(workbook, name: str, headings: tuple[str, ...], rows):
    ws = workbook.create_sheet(name)
    ws.append(headings)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headings))}1"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="203647")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for record in rows:
        ws.append([_safe_cell(value) for value in record])
    for index, heading in enumerate(headings, 1):
        width = 45 if any(term in heading.casefold() for term in (
            "source", "target", "proposal", "reason", "path",
        )) else 22
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def _bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _response(data: bytes, filename: str):
    return Response(
        content=data,
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": 'attachment; filename="' + filename + '"',
            "Cache-Control": "no-store",
        },
    )


def _batch_rows(job: dict):
    """One row per unique source offset, never infer old target proposals."""
    skipped = {
        item["offset"]: item for item in job.get("skipped", [])
    }
    recovered = job.get("recovered", {})
    details = job.get("verified_details", {})
    processed = job["next_offset"]
    for offset, source in enumerate(job["candidates"]):
        item = skipped.get(offset)
        fix = recovered.get(str(offset))
        confirmed = details.get(str(offset))
        reason = item.get("reason", "") if item else ""
        proposal = ""
        verified_rows = ""
        if fix is not None:
            status = "Recovered — verified in TIA"
            proposal = fix.get("proposal", "")
            verified_rows = fix.get("verified_rows", "")
        elif item:
            status = "Skipped — not imported"
            proposal = item.get("proposal", "")
        elif offset < processed:
            if confirmed:
                status = (
                    "Auto-repaired — verified in TIA"
                    if confirmed.get("auto_repaired")
                    else "Verified in TIA"
                )
                proposal = confirmed.get("proposal", "")
                verified_rows = confirmed.get("verified_rows", "")
            else:
                status = "Verified — legacy checkpoint"
                # Historical checkpoints retained a verified cursor but did
                # not save the exact proposal or per-source verified rows.
        else:
            status = (
                "Not processed — previous import state uncertain"
                if job["status"] in ("error", "interrupted", "processing")
                else "Not processed"
            )
        yield (
            offset + 1,
            status,
            source,
            proposal,
            reason,
            verified_rows,
            (confirmed or {}).get("batch_index", ""),
            fix.get("mode", "") if fix else job.get("model", ""),
        )


@router.get("/hmis/batch-translations/{job_id}/quality-report.xlsx")
def download_hmi_batch_quality_report(job_id: str):
    with _JOB_LOCK:
        job = _load_job(job_id)
        state = _public_job(job)
    book = Workbook()
    book.remove(book.active)
    metadata = [
        ("Job ID", job_id),
        ("HMI", state["hmi_name"]),
        ("ViewPath root", state["view_path_root"]),
        ("Source culture", job["request"]["source_culture"]),
        ("Target culture", job["request"]["target_culture"]),
        ("Model", job["model"]),
        ("Checkpoint state", state["status"]),
        ("Total unique sources", state["total_unique"]),
        ("Processed unique sources", state["processed_unique"]),
        ("Verified unique sources", state["completed_unique"]),
        ("Auto-repaired unique sources (recorded)",
         sum(bool(item.get("auto_repaired"))
             for item in job.get("verified_details", {}).values())),
        ("Total skipped attempts", state["skipped_total"]),
        ("Pending skipped unique sources", state["skipped_unique"]),
        ("Recovered skipped sources", state["recovered_unique"]),
        ("Confirmed TIA rows", state["completed_rows"]),
        ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Source", "Local checkpoint; NOT a new TIA inspection"),
        ("Legacy checkpoint caution",
         "Exact targets/auto-repair flags may be unknown for older verified entries."),
        ("Interrupted job caution",
         "Unprocessed rows may include an import with unconfirmed TIA state."),
    ]
    _sheet(book, "Summary", ("Metric", "Value"), metadata)
    _sheet(
        book, "Translations",
        ("Offset", "Status", "Source text", "Approved or rejected proposal",
         "Skip reason", "Verified TIA rows", "Batch", "Model or retry mode"),
        _batch_rows(job),
    )
    return _response(_bytes(book), "HMI_batch_" + job_id[:12] + "_quality.xlsx")


@router.post("/hmis/audit-existing-translations/quality-report.xlsx")
def download_hmi_existing_qa_report(request: TranslateHmiTextsRequest):
    # Re-export current TIA project texts and scan every eligible selected-HMI
    # row. The UI preview is intentionally capped; this report is not.
    audit = _audit_existing_hmi_translations(request, limit=1000000)
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Summary", ("Metric", "Value"), [
        ("HMI", audit.hmi_name),
        ("ViewPath root", request.view_path_root),
        ("Source culture", request.source_culture),
        ("Target culture", request.target_culture),
        ("Scanned rows", audit.scanned_rows),
        ("Filled existing targets", audit.filled_target_rows),
        ("Suspicious existing targets", audit.suspicious_rows),
        ("Generated UTC", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ("Audit method", "Read-only TIA export; deterministic heuristic warnings"),
        ("Scope", "Existing populated targets only; no modifications to TIA"),
        ("Caution", "Flags need manual evaluation; unflagged entries are not guaranteed correct"),
    ])
    _sheet(
        book, "Existing QA",
        ("Sheet", "Row", "ViewPath", "Source text", "Current target (TIA)", "QA reason"),
        (
            (item.sheet, item.row, item.view_path, item.source_text,
             item.target_text, "; ".join(item.reasons))
            for item in audit.findings
        ),
    )
    return _response(_bytes(book), "HMI_existing_QA_quality.xlsx")
