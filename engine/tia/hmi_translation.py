from pathlib import Path
import tempfile

from fastapi import APIRouter, HTTPException
from openpyxl import load_workbook

from engine.tia.adapter import read_project_saved
from engine.tia.compat import run_adapter
from engine.tia.ollama_translation import (
    resolve_ollama_model,
    translate_ollama_texts,
)
from engine.tia.models import (
    HmiTranslationExample,
    HmiProjectTextPreviewRow,
    HmiProjectTextRoot,
    HmiProjectTextSheetDiagnostic,
    HmiProjectTextsPreviewResponse,
    TranslateHmiTextsRequest,
    TranslateHmiTextsResponse,
)


router = APIRouter()


def _project_text_columns(worksheet, source_culture=None, target_culture=None):
    # Siemens TIA V21: Category | ViewPath | Internal ID |
    # Substitutions | en-US* | de-DE | ref="0" ...
    # Older/other layouts may use Category and Reference at the end.
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=50, values_only=True),
        start=1,
    ):
        names = {
            str(value).strip().casefold(): index
            for index, value in enumerate(row, start=1)
            if value is not None
        }
        category_col = names.get("category")
        reference_col = names.get("viewpath") or names.get("reference")
        if category_col is None or reference_col is None:
            continue

        if "viewpath" in names:
            source = (source_culture or "").strip().casefold()
            target = (target_culture or "").strip().casefold()
            source_col = names.get(source + "*") if source else None
            if source_col is None and source:
                source_col = names.get(source)
            if source_col is None:
                source_col = next(
                    (index for index, value in enumerate(row, start=1)
                     if isinstance(value, str) and value.strip().endswith("*")
                     and index > names["substitutions"]),
                    None,
                ) if "substitutions" in names else None

            target_col = names.get(target) if target else None
            if target_col is None and source:
                # In exports made before the target was enabled, TIA can
                # return source* and source as the two language columns.
                target_col = next(
                    (index for index, value in enumerate(row, start=1)
                     if index != source_col
                     and isinstance(value, str)
                     and value.strip().casefold() == source),
                    None,
                )
            if target_col is None and source_col is not None:
                # The next language column is adjacent to the source in
                # the two-language ExportProjectTexts XLSX format.
                possible = source_col + 1
                if possible <= len(row) and isinstance(row[possible - 1], str):
                    name = row[possible - 1].strip()
                    if name and not name.startswith('ref="'):
                        target_col = possible

            if source_col and target_col and source_col != target_col:
                return (
                    row_index, source_col, target_col,
                    category_col, reference_col,
                )
            continue

        # Previous workbook layout: languages precede Category, Reference.
        if category_col >= 3:
            return (
                row_index, category_col - 2, category_col - 1,
                category_col, reference_col,
            )
    return None



def _canonical_reference_part(
    value,
) -> str:
    text = str(
        value
        or
        ""
    ).casefold()

    result = []
    pending_separator = False

    for character in text:
        if character.isalnum():
            if pending_separator and result:
                result.append(
                    " "
                )

            result.append(
                character
            )

            pending_separator = False

        else:
            pending_separator = True

    return "".join(
        result
    ).strip()


def _segment_matches_target(
    segment: str,
    target: str,
) -> bool:
    if not target:
        return False

    if segment == target:
        return True

    if not segment.startswith(
        target
    ):
        return False

    if len(segment) == len(target):
        return True

    return not segment[
        len(target)
    ].isalnum()


def _view_path_root(reference) -> str:
    parts = str(reference or "").replace("/", "\\").split("\\")
    return parts[1].strip() if len(parts) >= 2 else ""


def _reference_matches_hmi(
    reference,
    device_name: str,
    hmi_name: str,
    view_path_root: str | None = None,
) -> bool:
    # Match the actual device/root of ViewPath, never arbitrary nested
    # screen or object names that might coincide with another HMI.
    root = _canonical_reference_part(_view_path_root(reference))
    targets = (
        [view_path_root]
        if view_path_root
        else [device_name, hmi_name]
    )
    return any(
        _segment_matches_target(root, _canonical_reference_part(name))
        for name in targets if name
    )



def _translate_texts(
    texts: list[str],
    target_culture: str,
    model_name: str | None = None,
    progress=None,
) -> dict[str, str]:
    """Translate exclusively through the local Ollama server."""
    return translate_ollama_texts(
        texts, target_culture, model_name, progress=progress
    )


def _preview_workbook(workbook, request):
    from collections import Counter

    matched_rows = []
    unmatched_rows = []
    references = []
    reference_set = set()
    roots = Counter()
    diagnostics = []
    total = 0
    matched = 0
    headers = ("", "")
    sheets_found = 0

    for worksheet in workbook.worksheets:
        if hasattr(worksheet, "reset_dimensions"):
            worksheet.reset_dimensions()

        # Capture raw structure even when no expected column headers exist.
        sample = []
        for index, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=5, values_only=True),
            start=1,
        ):
            sample.append(
                str(index) + ": " + " | ".join(
                    str(value or "")[:100] for value in row[:6]
                )
            )

        columns = _project_text_columns(
            worksheet, request.source_culture, request.target_culture,
        )
        diagnostics.append(
            HmiProjectTextSheetDiagnostic(
                sheet_name=worksheet.title,
                declared_rows=worksheet.max_row,
                declared_columns=worksheet.max_column,
                sample_rows=sample,
            )
        )
        if columns is None:
            continue

        sheets_found += 1
        header_row, source_col, target_col, category_col, reference_col = columns

        head = next(
            worksheet.iter_rows(
                min_row=header_row, max_row=header_row, values_only=True
            ), (),
        )
        if not headers[0]:
            headers = (
                str(head[source_col - 1] or "")
                if len(head) >= source_col else "",
                str(head[target_col - 1] or "")
                if len(head) >= target_col else "",
            )

        for index, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            def cell_value(column):
                return row[column - 1] if len(row) >= column else None

            source = cell_value(source_col)
            reference = str(cell_value(reference_col) or "").strip()
            category = str(cell_value(category_col) or "")
            if reference and len(references) < 50 and reference not in reference_set:
                reference_set.add(reference)
                references.append(reference)

            if source is None or not str(source).strip():
                continue

            total += 1
            if "viewpath" in str(head[reference_col - 1]).casefold():
                root = _view_path_root(reference)
                if root and category.casefold().startswith("<hmi"):
                    roots[root] += 1

            matches = _reference_matches_hmi(
                reference, request.device_name, request.hmi_name,
                request.view_path_root,
            )
            if matches:
                matched += 1

            dest = matched_rows if matches else unmatched_rows
            if len(dest) < 50:
                dest.append(
                    HmiProjectTextPreviewRow(
                        sheet=worksheet.title,
                        row=index,
                        source_text=str(source),
                        target_text=str(cell_value(target_col) or ""),
                        category=category,
                        reference=reference,
                        matches_hmi=matches,
                    )
                )

    return HmiProjectTextsPreviewResponse(
        hmi_name=request.hmi_name,
        device_name=request.device_name,
        source_culture=request.source_culture,
        target_culture=request.target_culture,
        source_column_header=headers[0],
        target_column_header=headers[1],
        worksheets_found=sheets_found,
        worksheet_count=len(workbook.worksheets),
        sheet_diagnostics=diagnostics,
        view_path_roots=[
            HmiProjectTextRoot(root=name, rows=count)
            for name, count in roots.most_common(50)
        ],
        total_text_rows=total,
        matched_text_rows=matched,
        reference_examples=references,
        sample_rows=matched_rows + unmatched_rows,
    )


@router.post(
    "/hmis/preview-texts",
    response_model=HmiProjectTextsPreviewResponse,
)
def preview_hmi_texts(request: TranslateHmiTextsRequest):
    if request.source_culture.casefold() == request.target_culture.casefold():
        raise HTTPException(400, "Source and target slots must differ.")

    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
            temporary_path = Path(handle.name)

        run_adapter(
            "export-project-texts",
            str(request.process_id),
            request.source_culture.strip(),
            request.target_culture.strip(),
            str(temporary_path),
        )
        if not temporary_path.exists() or temporary_path.stat().st_size == 0:
            raise HTTPException(502, "TIA Portal did not create the project-text XLSX file.")

        workbook = load_workbook(temporary_path, read_only=True, data_only=True)
        try:
            return _preview_workbook(workbook, request)
        finally:
            workbook.close()
    finally:
        if temporary_path and temporary_path.exists():
            temporary_path.unlink(missing_ok=True)


@router.post(
    "/hmis/translate-texts",
    response_model=TranslateHmiTextsResponse,
)
def translate_hmi_texts(
    request: TranslateHmiTextsRequest,
):
    source_culture = request.source_culture.strip()
    target_culture = request.target_culture.strip()

    if (
        source_culture.casefold()
        == target_culture.casefold()
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "The TIA source language slot and target language must differ. "
                "The actual language of the source text is detected from the text content."
            ),
        )

    # Resolve local model before changing any TIA project languages.
    model = resolve_ollama_model(request.ollama_model)

    run_adapter(
        "add-project-language",
        str(request.process_id),
        target_culture,
    )

    temporary_path = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=".xlsx",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)

        run_adapter(
            "export-project-texts",
            str(request.process_id),
            source_culture,
            target_culture,
            str(temporary_path),
        )

        if (
            not temporary_path.exists()
            or temporary_path.stat().st_size == 0
        ):
            raise HTTPException(
                status_code=502,
                detail=(
                    "TIA Portal did not create the project-text XLSX file."
                ),
            )

        workbook = load_workbook(temporary_path)
        matching_cells = []
        matched_row_indexes = {}
        seen_references = []

        for worksheet in workbook.worksheets:
            columns = _project_text_columns(
                worksheet,
                source_culture,
                target_culture,
            )

            if columns is None:
                continue

            (
                header_row,
                source_column,
                target_column,
                _category_column,
                reference_column,
            ) = columns

            # The actual V21 XLSX uses en-US* and <target culture>.
            # Before enabling the target culture, some exports instead
            # show en-US* / en-US. Retitle only that known placeholder;
            # reject any other ambiguous target column.
            reference_title = str(
                worksheet.cell(header_row, reference_column).value or ""
            ).strip().casefold()
            if reference_title == "viewpath":
                source_title = str(
                    worksheet.cell(header_row, source_column).value or ""
                ).strip()
                target_cell_header = worksheet.cell(header_row, target_column)
                actual_target_title = str(target_cell_header.value or "").strip()
                if source_title.rstrip("*").casefold() != source_culture.casefold():
                    raise HTTPException(
                        409,
                        "Unexpected TIA source language header: " + source_title,
                    )
                if actual_target_title.casefold() != target_culture.casefold():
                    if actual_target_title.casefold() != source_culture.casefold():
                        raise HTTPException(
                            409,
                            "Unexpected TIA target language header: "
                            + actual_target_title
                            + "; expected " + target_culture,
                        )
                    target_cell_header.value = target_culture

            matched_row_indexes[worksheet] = set()

            for row_index in range(
                header_row + 1,
                worksheet.max_row + 1,
            ):
                reference = worksheet.cell(
                    row=row_index,
                    column=reference_column,
                ).value

                if (
                    reference is not None
                    and str(reference).strip()
                    and len(seen_references) < 12
                ):
                    reference_text = str(
                        reference
                    ).strip()

                    if reference_text not in seen_references:
                        seen_references.append(
                            reference_text
                        )

                if not _reference_matches_hmi(
                    reference,
                    request.device_name,
                    request.hmi_name,
                    request.view_path_root,
                ):
                    continue

                source_cell = worksheet.cell(
                    row=row_index,
                    column=source_column,
                )
                target_cell = worksheet.cell(
                    row=row_index,
                    column=target_column,
                )
                source_text = source_cell.value

                if (
                    source_text is None
                    or not str(source_text).strip()
                ):
                    continue

                matched_row_indexes[worksheet].add(row_index)
                matching_cells.append(
                    (
                        str(source_text),
                        target_cell,
                    )
                )

        if not matching_cells:
            examples = (
                " | ".join(
                    seen_references[:5]
                )
                if seen_references
                else "(no non-empty Reference values found)"
            )

            raise HTTPException(
                status_code=404,
                detail=(
                    "No project-text rows were found for HMI "
                    f'"{request.hmi_name}". '
                    "Example Reference values from the exported XLSX: "
                    + examples
                ),
            )

        # For a pilot, import at most max_rows distinct meaningful source
        # texts. Avoid wasting test calls on duplicate labels or bare symbols.
        if request.max_rows is not None:
            pilot = []
            pilot_texts = set()
            for text, cell in matching_cells:
                if text in pilot_texts or not any(ch.isalpha() for ch in text):
                    continue
                pilot.append((text, cell))
                pilot_texts.add(text)
                if len(pilot) >= request.max_rows:
                    break

            if not pilot:
                raise HTTPException(
                    422,
                    "No textual HMI rows are available for the pilot import.",
                )

            matching_cells = pilot
            for worksheet in matched_row_indexes:
                matched_row_indexes[worksheet] = set()
            for _source, target_cell in matching_cells:
                matched_row_indexes[target_cell.parent].add(target_cell.row)

        unique_texts = list(
            dict.fromkeys(
                text
                for text, _cell in matching_cells
            )
        )

        translations = _translate_texts(
            unique_texts,
            target_culture,
            model,
        )

        translated_examples = []
        for source_text, target_cell in matching_cells[:10]:
            columns = _project_text_columns(
                target_cell.parent, source_culture, target_culture,
            )
            reference_column = columns[4] if columns else None
            ref_text = (
                str(target_cell.parent.cell(
                    target_cell.row, reference_column
                ).value or "")
                if reference_column else ""
            )
            translated_examples.append(
                HmiTranslationExample(
                    view_path=ref_text,
                    source_text=source_text,
                    target_text=translations[source_text],
                )
            )

        for source_text, target_cell in matching_cells:
            target_cell.value = translations[
                source_text
            ]

        # Import a workbook containing ONLY selected HMI text rows.
        # Otherwise renaming a placeholder language column could cause
        # unrelated HMI texts to be imported as translations as well.
        for worksheet, keep_rows in matched_row_indexes.items():
            columns = _project_text_columns(
                worksheet, source_culture, target_culture,
            )
            if columns is None:
                continue
            header_row = columns[0]
            excluded_runs = []
            run_start = None
            for index in range(header_row + 1, worksheet.max_row + 1):
                if index in keep_rows:
                    if run_start is not None:
                        excluded_runs.append((run_start, index - 1))
                        run_start = None
                elif run_start is None:
                    run_start = index
            if run_start is not None:
                excluded_runs.append((run_start, worksheet.max_row))
            for start, end in reversed(excluded_runs):
                worksheet.delete_rows(start, end - start + 1)

        workbook.save(temporary_path)

        import_lines = run_adapter(
            "import-project-texts",
            str(request.process_id),
            str(temporary_path),
            "false",
        )

        import_state = ""
        import_log = ""

        for line in import_lines:
            if not line.startswith(
                "PROJECT_TEXTS_IMPORTED|"
            ):
                continue

            parts = line.split("|", 2)

            if len(parts) >= 2:
                import_state = parts[1]

            if len(parts) == 3:
                import_log = parts[2]

        # ProjectTextResult.State is authoritative for import success.
        # Do not report success merely because Openness returned a result.
        if import_state.strip().casefold() not in {"success", "succeeded"}:
            raise HTTPException(
                status_code=502,
                detail=(
                    "TIA project texts import did not report success. "
                    "Import state: " + (import_state or "(missing)") + ". "
                    "Import log: " + (import_log or "(not returned)") + ". "
                    "Check the log before trying another import."
                ),
            )

        return TranslateHmiTextsResponse(
            status="translated",
            hmi_name=request.hmi_name,
            source_culture=source_culture,
            target_culture=target_culture,
            translated_rows=len(
                matching_cells
            ),
            unique_texts=len(
                unique_texts
            ),
            translated_examples=translated_examples,
            import_state=import_state,
            import_log=import_log,
            project_saved=read_project_saved(
                import_lines
            ),
        )

    finally:
        if (
            temporary_path
            and temporary_path.exists()
        ):
            temporary_path.unlink(
                missing_ok=True
            )
