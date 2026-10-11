from pathlib import Path
import re
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import engine.tia_openness as tia_openness
from engine.main import app


client = TestClient(app)


def _write_sample_export(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["English (United States)", "Polish (Poland)", "Category", "Reference"])
    sheet.append(["Zmiana", None, "Screen", "Project\\HMI_RT_1 [Runtime]\\Main"])
    sheet.append(["Zgniatarka", None, "Screen", "Project\\HMI_RT_10 [Runtime]\\Main"])
    sheet.append(["Zezwolenie na jazdę", None, "Screen", "Project\\HM_Conveyors [HMI]\\Main"])
    workbook.save(path)


def test_preview_is_read_only_and_filters_exact_hmi(monkeypatch):
    commands = []

    def adapter(*args):
        commands.append(args)
        if args[0] == "export-project-texts":
            _write_sample_export(Path(args[4]))
            return ["PROJECT_TEXTS_EXPORTED|" + args[4]]
        raise AssertionError("Preview must not add languages, import or save")

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)

    response = client.post(
        "/api/v1/tia/hmis/preview-texts",
        json={
            "process_id": 3544,
            "device_name": "HMI_RT_1",
            "hmi_name": "HMI_RT_1",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
        },
    )

    assert response.status_code == 200
    result = response.json()
    assert result["total_text_rows"] == 3
    assert result["matched_text_rows"] == 1
    assert result["source_column_header"] == "English (United States)"
    assert result["target_column_header"] == "Polish (Poland)"
    assert any(r["source_text"] == "Zmiana" and r["matches_hmi"]
               for r in result["sample_rows"])
    assert any("HMI_RT_10" in ref for ref in result["reference_examples"])
    assert len(commands) == 1
    assert commands[0][0] == "export-project-texts"


def test_preview_rejects_same_language_without_export(monkeypatch):
    def forbidden(*args):
        raise AssertionError("No adapter call expected")

    monkeypatch.setattr(tia_openness, "_run_adapter", forbidden)
    response = client.post(
        "/api/v1/tia/hmis/preview-texts",
        json={
            "process_id": 3544,
            "device_name": "HMI_RT_1",
            "hmi_name": "HMI_RT_1",
            "source_culture": "pl-PL",
            "target_culture": "pl-PL",
        },
    )
    assert response.status_code == 400


def test_preview_handles_read_only_sheet_with_unknown_dimensions(monkeypatch):
    def adapter(*args):
        assert args[0] == "export-project-texts"
        path = Path(args[-1])
        _write_sample_export(path)

        # Simulate an XLSX exported without worksheet dimension metadata.
        with ZipFile(path, "r") as source:
            members = [
                (item, source.read(item.filename))
                for item in source.infolist()
            ]
        with ZipFile(path, "w") as destination:
            for item, data in members:
                if item.filename == "xl/worksheets/sheet1.xml":
                    data, count = re.subn(
                        rb"<dimension\b[^>]*/>", b"", data, count=1
                    )
                    assert count == 1
                destination.writestr(item, data)

        workbook = load_workbook(path, read_only=True)
        try:
            assert workbook.active.max_row is None
        finally:
            workbook.close()
        return ["PROJECT_TEXTS_EXPORTED|" + str(path)]

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)

    response = client.post(
        "/api/v1/tia/hmis/preview-texts",
        json={
            "process_id": 13940,
            "device_name": "HMI_RT_1",
            "hmi_name": "HMI_RT_1",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total_text_rows"] == 3
    assert data["matched_text_rows"] == 1
    assert data["sample_rows"][0]["source_text"] == "Zmiana"


def _write_tia_v21_export(path: Path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "User Texts"
    sheet.append([
        "Category", "ViewPath", "Internal ID", "Substitutions",
        "en-US*", "en-US", 'ref="0"',
    ])
    sheet.append([
        "<HMI screen>", "Project\\HMI_00_Main\\Screens\\Main\\Text",
        "opaque-1", "", "Zmiana", "Zmiana", "",
    ])
    sheet.append([
        "<Alarm>", "Project\\HMI_00_Main\\Alarms\\Alarm_1",
        "opaque-2", "", "Zgniatarka", "Zgniatarka", "",
    ])
    sheet.append([
        "<HMI screen>", "Project\\HMI_Other\\Screens\\Main\\Text",
        "opaque-3", "", "Other HMI text", "Other HMI text", "",
    ])
    system = workbook.create_sheet("System Texts")
    system.append([
        "Category", "ViewPath", "Internal ID", "Substitutions",
        "en-US*", "en-US", 'ref="0"',
    ])
    workbook.save(path)


def test_preview_real_tia_v21_layout_and_explicit_view_path_root(monkeypatch):
    def adapter(*args):
        assert args[0] == "export-project-texts"
        _write_tia_v21_export(Path(args[-1]))
        return ["PROJECT_TEXTS_EXPORTED|" + args[-1]]

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    body = {
        "process_id": 13940,
        "device_name": "Unknown_inventory_device",
        "hmi_name": "HMI_RT_1",
        "source_culture": "en-US",
        "target_culture": "pl-PL",
    }
    response = client.post("/api/v1/tia/hmis/preview-texts", json=body)
    assert response.status_code == 200
    info = response.json()
    assert info["worksheets_found"] == 2
    assert info["worksheet_count"] == 2
    assert info["total_text_rows"] == 3
    assert info["matched_text_rows"] == 0
    assert info["source_column_header"] == "en-US*"
    assert info["target_column_header"] == "en-US"
    assert any(x["root"] == "HMI_00_Main" for x in info["view_path_roots"])

    body["view_path_root"] = "HMI_00_Main"
    response = client.post("/api/v1/tia/hmis/preview-texts", json=body)
    assert response.status_code == 200
    assert response.json()["matched_text_rows"] == 2


def test_translate_imports_only_selected_hmi_and_renames_target_header(
    monkeypatch,
):
    import engine.tia.hmi_translation as translation

    commands = []
    imported = []

    def adapter(*args):
        commands.append(args[0])
        if args[0] == "add-project-language":
            return ["LANGUAGE_ADDED|de-DE"]
        if args[0] == "export-project-texts":
            _write_tia_v21_export(Path(args[-1]))
            return ["PROJECT_TEXTS_EXPORTED|" + args[-1]]
        if args[0] == "import-project-texts":
            assert args[-1] == "false"
            file_path = Path(args[2])
            workbook = load_workbook(file_path)
            rows = list(workbook["User Texts"].values)
            assert rows[0][4] == "en-US*"
            assert rows[0][5] == "de-DE"
            assert len(rows) == 3
            assert all("HMI_00_Main" in row[1] for row in rows[1:])
            assert rows[1][2] == "opaque-1"
            assert rows[2][2] == "opaque-2"
            assert rows[1][5] == "Shift"
            assert rows[2][5] == "Crusher"
            imported.append(True)
            workbook.close()
            return ["PROJECT_TEXTS_IMPORTED|Success|", "PROJECT_SAVED|Project"]
        raise AssertionError("Unexpected adapter command " + args[0])

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    monkeypatch.setattr(
        translation,
        "_translate_texts",
        lambda texts, target: {
            "Zmiana": "Shift", "Zgniatarka": "Crusher",
        },
    )
    response = client.post(
        "/api/v1/tia/hmis/translate-texts",
        json={
            "process_id": 13940,
            "device_name": "Unknown_inventory_device",
            "hmi_name": "HMI_RT_1",
            "view_path_root": "HMI_00_Main",
            "source_culture": "en-US",
            "target_culture": "de-DE",
        },
    )
    assert response.status_code == 200
    assert imported
    assert response.json()["translated_rows"] == 2
    assert commands == [
        "add-project-language",
        "export-project-texts",
        "import-project-texts",
    ]


def test_view_path_match_is_limited_to_device_root():
    from engine.tia.hmi_translation import _reference_matches_hmi

    assert not _reference_matches_hmi(
        "Project\\UnrelatedDevice\\Screens\\HMI_RT_1\\Text",
        "HMI_RT_1",
        "HMI_RT_1",
    )
    assert _reference_matches_hmi(
        "Project\\HMI_00_Main\\Screens\\HMI_RT_1\\Text",
        "HMI_RT_1",
        "HMI_RT_1",
        "HMI_00_Main",
    )


def test_translation_returns_error_when_tia_import_reports_failure(monkeypatch):
    import engine.tia.hmi_translation as translation

    def adapter(*args):
        if args[0] == "add-project-language":
            return ["LANGUAGE_ADDED|pl-PL"]
        if args[0] == "export-project-texts":
            _write_tia_v21_export(Path(args[-1]))
            return ["PROJECT_TEXTS_EXPORTED|" + args[-1]]
        if args[0] == "import-project-texts":
            return [
                "PROJECT_TEXTS_IMPORTED|Failed|C:\\\\TIA\\\\import-log.txt",
                "PROJECT_SAVED|Project",
            ]
        raise AssertionError(args[0])

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    monkeypatch.setattr(
        translation,
        "_translate_texts",
        lambda texts, target: {text: text for text in texts},
    )
    response = client.post(
        "/api/v1/tia/hmis/translate-texts",
        json={
            "process_id": 13940,
            "device_name": "NotMatching",
            "hmi_name": "HMI_RT_1",
            "view_path_root": "HMI_00_Main",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
            "max_rows": 10,
        },
    )
    assert response.status_code == 502
    assert "Import state: Failed" in response.json()["detail"]
    assert "import-log.txt" in response.json()["detail"]


def test_pilot_limits_number_of_unique_rows(monkeypatch):
    import engine.tia.hmi_translation as translation

    translated = []
    imported = []

    def adapter(*args):
        if args[0] == "add-project-language":
            return ["LANGUAGE_ADDED|pl-PL"]
        if args[0] == "export-project-texts":
            _write_tia_v21_export(Path(args[-1]))
            return ["PROJECT_TEXTS_EXPORTED|" + args[-1]]
        if args[0] == "import-project-texts":
            workbook = load_workbook(Path(args[2]))
            rows = list(workbook["User Texts"].values)
            imported.extend(rows[1:])
            workbook.close()
            return ["PROJECT_TEXTS_IMPORTED|Success|", "PROJECT_SAVED|Project"]
        raise AssertionError(args[0])

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    monkeypatch.setattr(
        translation,
        "_translate_texts",
        lambda texts, target: (
            translated.extend(texts) or {text: text for text in texts}
        ),
    )
    response = client.post(
        "/api/v1/tia/hmis/translate-texts",
        json={
            "process_id": 13940,
            "device_name": "Unrelated inventory",
            "hmi_name": "HMI_RT_1",
            "view_path_root": "HMI_00_Main",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
            "max_rows": 1,
        },
    )
    assert response.status_code == 200
    assert response.json()["translated_rows"] == 1
    assert len(imported) == 1
    assert len(translated) == 1
    assert "HMI_00_Main" in imported[0][1]
