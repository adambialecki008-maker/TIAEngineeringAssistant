"""Integration tests for sequential, checkpointed HMI batch imports."""
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from engine.main import app
from engine.tia import hmi_batches as batches
from engine.tia import hmi_review as review
import engine.tia_openness as tia_openness


client = TestClient(app)
REQUEST = {
    "process_id": 13940,
    "device_name": "HMI_RT_1",
    "hmi_name": "HMI_RT_1",
    "view_path_root": "HMI_00_Main",
    "source_culture": "en-US",
    "target_culture": "pl-PL",
    "ollama_model": "qwen2.5:3b",
    "batch_size": 1,
    "total_to_translate": 2,
}


def setup_batch_mock(monkeypatch, tmp_path, *, fail_on_import=None,
                     first_source="Zgniatarka"):
    monkeypatch.setattr(batches, "JOB_DIR", tmp_path)
    monkeypatch.setattr(
        batches, "resolve_ollama_model",
        lambda _requested: "qwen2.5:3b",
    )
    batches._ACTIVE_JOBS.clear()
    review._REVIEWS.clear()
    live = {"main-1": "Text", "main-2": "Text", "other-1": "Original"}
    commands = []
    imports = []

    def export_file(path):
        wb = Workbook()
        sheet = wb.active
        sheet.title = "User Texts"
        sheet.append([
            "Category", "ViewPath", "Internal ID", "Substitutions",
            "en-US*", "pl-PL", 'ref="0"',
        ])
        for row_id, root, source in [
            ("main-1", "HMI_00_Main", first_source),
            ("main-2", "HMI_00_Main", "Za zaworem"),
            ("other-1", "Another_HMI", "Unrelated"),
        ]:
            sheet.append([
                "<HMI screen>", f"Project\\{root}\\Screens\\Main\\{row_id}",
                row_id, "", source, live[row_id], "",
            ])
        other = wb.create_sheet("System Texts")
        other.append([
            "Category", "ViewPath", "Internal ID", "Substitutions",
            "en-US*", "pl-PL", 'ref="0"',
        ])
        wb.save(path)
        wb.close()

    def adapter(*args):
        commands.append(args[0])
        if args[0] == "export-project-texts":
            export_file(Path(args[-1]))
            return ["PROJECT_TEXTS_EXPORTED|" + args[-1]]
        if args[0] == "add-project-language":
            return ["LANGUAGE_ADDED|pl-PL"]
        if args[0] == "import-project-texts":
            if fail_on_import and len(imports) + 1 == fail_on_import:
                return ["PROJECT_TEXTS_IMPORTED|Failed|test import failed"]
            wb = load_workbook(args[2])
            vals = [
                (row[2], row[5])
                for row in list(wb["User Texts"].values)[1:]
            ]
            # Existing text values are retained for all sibling rows,
            # including those belonging to a later batch.
            imports.append(vals)
            for row_id, target in vals:
                live[row_id] = target
            wb.close()
            return [
                "PROJECT_TEXTS_IMPORTED|Success|test-import.log",
                "PROJECT_SAVED|SampleProject",
            ]
        raise AssertionError(args)

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            source: "Przetłumaczone " + source for source in texts
        },
    )
    return live, commands, imports


def test_two_batches_import_each_to_tia_and_save_checkpoint(monkeypatch, tmp_path):
    live, commands, imports = setup_batch_mock(monkeypatch, tmp_path)

    started = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json=REQUEST,
    )
    assert started.status_code == 200, started.text
    job = started.json()
    assert job["total_unique"] == 2
    assert job["completed_unique"] == 0
    assert job["status"] == "ready"
    job_id = job["job_id"]

    first = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert first.status_code == 200, first.text
    first_result = first.json()
    assert first_result["completed_unique"] == 1
    assert first_result["completed_rows"] == 1
    assert first_result["last_verified_batch"]["verified_rows"] == 1
    assert first_result["status"] == "ready"
    assert len(imports) == 1
    assert live["main-1"] == "Przetłumaczone Zgniatarka"
    assert live["main-2"] == "Text"
    assert live["other-1"] == "Original"

    # New endpoint instance reads checkpoint from disk, not just memory.
    assert (tmp_path / f"{job_id}.json").exists()
    checkpoint = client.get(f"/api/v1/tia/hmis/batch-translations/{job_id}")
    assert checkpoint.json()["completed_unique"] == 1

    second = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert second.status_code == 200, second.text
    second_result = second.json()
    assert second_result["completed_unique"] == 2
    assert second_result["completed_rows"] == 2
    assert second_result["status"] == "completed"
    assert len(imports) == 2
    assert live["main-1"] == "Przetłumaczone Zgniatarka"
    assert live["main-2"] == "Przetłumaczone Za zaworem"
    assert live["other-1"] == "Original"
    assert commands.count("import-project-texts") == 2
    assert commands.count("export-project-texts") == 7


def test_failed_batch_stops_without_advancing_verified_cursor(monkeypatch, tmp_path):
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, fail_on_import=2
    )
    start = client.post(
        "/api/v1/tia/hmis/batch-translations/start", json=REQUEST,
    ).json()
    job_id = start["job_id"]
    first = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert first.status_code == 200
    second = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert second.status_code == 502
    job = client.get(
        f"/api/v1/tia/hmis/batch-translations/{job_id}",
    ).json()
    assert job["status"] == "error"
    assert job["completed_unique"] == 1
    assert job["completed_rows"] == 1
    assert len(imports) == 1
    assert live["main-1"] == "Przetłumaczone Zgniatarka"
    assert live["main-2"] == "Text"
    again = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert again.status_code == 409
    assert commands.count("import-project-texts") == 2


def test_pause_then_resume_skips_verified_batches(monkeypatch, tmp_path):
    live, _, imports = setup_batch_mock(monkeypatch, tmp_path)
    start = client.post(
        "/api/v1/tia/hmis/batch-translations/start", json=REQUEST,
    ).json()
    job_id = start["job_id"]
    first = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert first.json()["completed_unique"] == 1
    paused = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/pause",
    )
    assert paused.status_code == 200
    assert paused.json()["status"] == "paused"
    resumed = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert resumed.status_code == 200, resumed.text
    assert resumed.json()["status"] == "completed"
    assert resumed.json()["completed_unique"] == 2
    assert len(imports) == 2


def test_explicit_view_path_root_required(monkeypatch, tmp_path):
    setup_batch_mock(monkeypatch, tmp_path)
    response = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "view_path_root": None},
    )
    assert response.status_code == 422
    assert not list(tmp_path.glob("*.json"))


def test_batch_size_greater_than_total_imports_only_total(monkeypatch, tmp_path):
    live, commands, imports = setup_batch_mock(monkeypatch, tmp_path)
    response = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 5, "total_to_translate": 2},
    )
    assert response.status_code == 200, response.text
    job = response.json()
    assert job["total_batches"] == 1
    assert job["total_unique"] == 2

    completed = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job['job_id']}/next", json={},
    )
    assert completed.status_code == 200, completed.text
    result = completed.json()
    assert result["status"] == "completed"
    assert result["completed_unique"] == 2
    assert result["last_verified_batch"]["verified_unique"] == 2
    assert len(imports) == 1
    assert commands.count("import-project-texts") == 1


def test_batch_source_copy_does_not_call_ollama(monkeypatch, tmp_path):
    live, commands, imports = setup_batch_mock(monkeypatch, tmp_path)
    def forbidden_translation(*args, **kwargs):
        raise AssertionError("Ollama may not be invoked in exact copy mode")
    monkeypatch.setattr(review, "_translate_texts", forbidden_translation)
    monkeypatch.setattr(batches, "resolve_ollama_model", forbidden_translation)
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1,
              "translation_strategy": "copy"},
    )
    assert created.status_code == 200, created.text
    response = client.post(
        f"/api/v1/tia/hmis/batch-translations/{created.json()['job_id']}/next",
        json={},
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "completed"
    assert live["main-1"] == "Zgniatarka"
    assert "import-project-texts" in commands


def test_numbered_polish_instruction_imports_after_identifier_guard_fix(
    monkeypatch, tmp_path,
):
    """Check the full Ollama -> XLSX -> TIA re-export path for '1.Naciagnij'."""
    source = (
        "1.Naciagnij linke \n"
        "do tej pozycji \n"
        "(koncowka w obrebie\n zielonego wskaznika)"
    )
    target = (
        "1.Naciągnij linkę\n"
        "do tej pozycji\n"
        "(końcówka w obrębie\nzielonego wskaźnika)"
    )
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, culture, model, progress=None: {
            text: target if text == source else text for text in texts
        },
    )
    started = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    assert started.status_code == 200, started.text
    run = client.post(
        f"/api/v1/tia/hmis/batch-translations/{started.json()['job_id']}/next",
        json={},
    )
    assert run.status_code == 200, run.text
    assert run.json()["completed_unique"] == 1
    assert run.json()["last_verified_batch"]["verified_rows"] == 1
    assert run.json()["status"] == "completed"
    assert live["main-1"] == target
    assert commands.count("import-project-texts") == 1
    assert commands.count("export-project-texts") == 4



def test_unsafe_ollama_proposal_is_skipped_but_valid_sibling_is_imported(
    monkeypatch, tmp_path,
):
    source = (
        "Dalmierz z intuicyjna nastawa punktu przelaczania\n"
        "OID251 IFM\nDomyslna nastawa powinna byc ok 170cm d"
    )
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    unsafe = "(Pneumatyczna) " + source[:-9]
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, culture, model, progress=None: {
            text: unsafe if text == source else "Za zaworem" for text in texts
        },
    )
    started = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 2},
    )
    assert started.status_code == 200, started.text
    job_id = started.json()["job_id"]
    response = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert response.status_code == 200, response.text
    job = response.json()
    assert job["status"] == "completed"
    assert job["processed_unique"] == 2
    assert job["completed_unique"] == 1
    assert job["skipped_unique"] == 1
    assert job["last_batch"]["skipped_unique"] == 1
    assert job["last_verified_batch"]["verified_unique"] == 1
    assert job["skipped_examples"][0]["source"] == source[:160]
    assert "changed count of '('" in job["skipped_examples"][0]["reason"]
    assert live["main-1"] == "Text"
    assert live["main-2"] == "Za zaworem"
    assert len(imports) == 1
    assert commands.count("import-project-texts") == 1

    persisted = client.get(
        f"/api/v1/tia/hmis/batch-translations/{job_id}"
    ).json()
    assert persisted["skipped_unique"] == 1
    assert persisted["completed_unique"] == 1


def test_whole_unsafe_batch_is_skipped_and_later_batch_continues(
    monkeypatch, tmp_path,
):
    source = "(Hydrauliczna)"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, culture, model, progress=None: {
            text: "(cykliski" if text == source else "Za zaworem"
            for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start", json=REQUEST
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job_id"]
    first = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert first.status_code == 200, first.text
    state = first.json()
    assert state["status"] == "ready"
    assert state["processed_unique"] == 1
    assert state["completed_unique"] == 0
    assert state["skipped_unique"] == 1
    assert state["last_verified_batch"] is None
    assert len(imports) == 0
    assert "import-project-texts" not in commands
    assert live["main-1"] == "Text"

    next_batch = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert next_batch.status_code == 200, next_batch.text
    finished = next_batch.json()
    assert finished["status"] == "completed"
    assert finished["processed_unique"] == 2
    assert finished["completed_unique"] == 1
    assert finished["skipped_unique"] == 1
    assert finished["last_verified_batch"]["verified_unique"] == 1
    assert len(imports) == 1
    assert live["main-1"] == "Text"
    assert live["main-2"] == "Za zaworem"


def test_entire_job_can_finish_with_only_skips_without_import(monkeypatch, tmp_path):
    source = "(Hydrauliczna)"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, culture, model, progress=None: {
            text: "({Hydrauliczna})" for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "total_to_translate": 1},
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job_id"]
    finished = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert finished.status_code == 200, finished.text
    data = finished.json()
    assert data["status"] == "completed"
    assert data["completed_unique"] == 0
    assert data["processed_unique"] == 1
    assert data["skipped_unique"] == 1
    assert data["last_verified_batch"] is None
    assert not imports
    assert "import-project-texts" not in commands
    assert live["main-1"] == "Text"



def test_batch_normalizes_literal_ollama_newline_before_import(
    monkeypatch, tmp_path,
):
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source="Awaria PLC",
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: r"Awaria \nPLC" for text in texts
        },
    )
    started = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    assert started.status_code == 200, started.text
    response = client.post(
        f"/api/v1/tia/hmis/batch-translations/{started.json()['job_id']}/next",
        json={},
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "completed"
    assert response.json()["skipped_unique"] == 0
    assert response.json()["completed_unique"] == 1
    assert live["main-1"] == "Awaria PLC"
    assert "import-project-texts" in commands
    assert len(imports) == 1



def test_batch_repairs_literal_newline_already_present_in_source(
    monkeypatch, tmp_path,
):
    source = r"Awaria \nPLC"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: text for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    assert created.status_code == 200, created.text
    finished = client.post(
        f"/api/v1/tia/hmis/batch-translations/{created.json()['job_id']}/next",
        json={},
    )
    assert finished.status_code == 200, finished.text
    assert finished.json()["status"] == "completed"
    assert finished.json()["skipped_unique"] == 0
    # The literal source escape is an intended line break (TIA shows ¶).
    assert live["main-1"] == "Awaria\nPLC"
    assert len(imports) == 1
    assert "import-project-texts" in commands



def test_batch_keeps_actual_tia_newline_when_source_has_one(
    monkeypatch, tmp_path,
):
    source = "Awaria\nhardware"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: r"Awaria\nhardware" for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    assert created.status_code == 200, created.text
    finished = client.post(
        f"/api/v1/tia/hmis/batch-translations/{created.json()['job_id']}/next",
        json={},
    )
    assert finished.status_code == 200, finished.text
    assert finished.json()["status"] == "completed"
    assert finished.json()["skipped_unique"] == 0
    assert live["main-1"] == "Awaria\nhardware"
    assert len(imports) == 1
    assert "import-project-texts" in commands



def test_skipped_manager_manual_prepare_then_verified_import(monkeypatch, tmp_path):
    source = "Pneumatyczny"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(Pneumatyczny" for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    assert created.status_code == 200, created.text
    job_id = created.json()["job_id"]
    ran = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    )
    assert ran.status_code == 200, ran.text
    assert ran.json()["status"] == "completed"
    assert ran.json()["skipped_unique"] == 1
    assert not imports

    skipped = client.get(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped"
    )
    assert skipped.status_code == 200, skipped.text
    assert skipped.json()["items"][0]["source"] == source
    assert skipped.json()["items"][0]["rejected_proposal"] == "(Pneumatyczny"
    assert skipped.json()["items"][0]["resolved"] is False

    prepared = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/prepare",
        json={"mode": "manual", "proposal": source},
    )
    assert prepared.status_code == 200, prepared.text
    body = prepared.json()
    assert body["proposal"] == source
    assert body["rows"] == 1
    assert live["main-1"] == "Text"
    assert not imports

    approved = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/import",
        json={"review_id": body["review_id"]},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "verified"
    assert approved.json()["verified_rows"] == 1
    assert live["main-1"] == source
    assert live["main-2"] == "Text"
    assert live["other-1"] == "Original"
    assert len(imports) == 1
    assert commands.count("import-project-texts") == 1

    state = client.get(f"/api/v1/tia/hmis/batch-translations/{job_id}").json()
    assert state["completed_unique"] == 1
    assert state["skipped_total"] == 1
    assert state["skipped_unique"] == 0
    assert state["recovered_unique"] == 1
    assert state["completed_rows"] == 1
    assert client.get(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped"
    ).json()["items"][0]["resolved"] is True
    repeat = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/import",
        json={"review_id": body["review_id"]},
    )
    assert repeat.status_code == 409
    assert len(imports) == 1


def test_skipped_retry_ollama_must_be_reviewed_before_import(monkeypatch, tmp_path):
    source = "(Hydrauliczna)"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(Hydrauliczna" for text in texts
        },
    )
    job_id = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    ).json()["job_id"]
    assert client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    ).status_code == 200

    unsafe_manual = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/prepare",
        json={"mode": "manual", "proposal": "(Hydrauliczna"},
    )
    assert unsafe_manual.status_code == 502
    assert not imports
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(Hydrauliczna)" for text in texts
        },
    )
    prepared = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/prepare",
        json={"mode": "ollama", "ollama_model": "qwen2.5:3b"},
    )
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()["model"] == "qwen2.5:3b"
    assert live["main-1"] == "Text"
    assert len(imports) == 0
    approved = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/import",
        json={"review_id": prepared.json()["review_id"]},
    )
    assert approved.status_code == 200, approved.text
    assert live["main-1"] == source
    assert len(imports) == 1
    assert commands.count("import-project-texts") == 1


def test_skipped_recovery_requires_completed_job_and_matching_review(
    monkeypatch, tmp_path,
):
    source = "Pneumatyczny"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(Pneumatyczny" for text in texts
        },
    )
    job_id = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 2},
    ).json()["job_id"]
    assert client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    ).status_code == 200
    forbidden = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/prepare",
        json={"mode": "manual", "proposal": source},
    )
    assert forbidden.status_code == 409
    assert live["main-1"] == "Text"

    # Finalize the remaining batch, then refuse a token not bound to this skip.
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: text for text in texts
        },
    )
    assert client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next", json={},
    ).status_code == 200
    invalid_token = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/import",
        json={"review_id": "a" * 32},
    )
    assert invalid_token.status_code == 409
    assert live["main-1"] == "Text"
    assert commands.count("import-project-texts") == 1



def test_skipped_review_cannot_bypass_checkpoint_via_generic_import(
    monkeypatch, tmp_path,
):
    source = "Pneumatyczny"
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(Pneumatyczny" for text in texts
        },
    )
    job_id = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    ).json()["job_id"]
    assert client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next",
        json={},
    ).status_code == 200
    prepared = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/prepare",
        json={"mode": "manual", "proposal": source},
    ).json()

    bypass = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert bypass.status_code == 409, bypass.text
    assert not imports
    assert live["main-1"] == "Text"

    valid = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped/0/import",
        json={"review_id": prepared["review_id"]},
    )
    assert valid.status_code == 200, valid.text
    assert valid.json()["job"]["recovered_unique"] == 1
    assert len(imports) == 1


def test_skipped_list_returns_original_full_source_from_checkpoint(
    monkeypatch, tmp_path,
):
    source = ("Pneumatyczny " * 22).strip()
    live, commands, imports = setup_batch_mock(
        monkeypatch, tmp_path, first_source=source,
    )
    monkeypatch.setattr(
        review, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: "(" + text for text in texts
        },
    )
    created = client.post(
        "/api/v1/tia/hmis/batch-translations/start",
        json={**REQUEST, "batch_size": 1, "total_to_translate": 1},
    )
    job_id = created.json()["job_id"]
    run = client.post(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/next",
        json={},
    )
    assert run.status_code == 200, run.text
    response = client.get(
        f"/api/v1/tia/hmis/batch-translations/{job_id}/skipped"
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["items"]) == 1
    assert response.json()["items"][0]["source"] == source
    assert len(source) > 160
    assert response.json()["items"][0]["resolved"] is False
    assert not imports
