"""Staged HMI translation: no TIA changes before approval and verified import."""
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import engine.tia_openness as tia_openness
from engine.main import app
import engine.tia.hmi_review as review_module

client = TestClient(app)


def _export_file(path: Path, live_targets: dict):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "User Texts"
    sheet.append([
        "Category", "ViewPath", "Internal ID",
        "Substitutions", "en-US*", "pl-PL", 'ref="0"',
    ])
    records = [
        ("main-1", "Project\\HMI_00_Main\\Screens\\Main\\Text",
         "212BP2 Za zaworem", "Text"),
        ("main-2", "Project\\HMI_00_Main\\Screens\\Status\\Text",
         "Zgniatarka", "Text"),
        ("other-1", "Project\\Another_HMI\\Screens\\Main\\Text",
         "Do not translate", "Do not translate"),
    ]
    for internal_id, path_value, source, default_target in records:
        sheet.append([
            "<HMI screen>", path_value, internal_id, "",
            source, live_targets.get(internal_id, default_target), "",
        ])
    sheet2 = workbook.create_sheet("System Texts")
    sheet2.append([
        "Category", "ViewPath", "Internal ID",
        "Substitutions", "en-US*", "pl-PL", 'ref="0"',
    ])
    workbook.save(path)


REQUEST = {
    "process_id": 13940,
    "device_name": "HMI_RT_1",
    "hmi_name": "HMI_RT_1",
    "view_path_root": "HMI_00_Main",
    "source_culture": "en-US",
    "target_culture": "pl-PL",
    "max_rows": 10,
}


def _setup(monkeypatch, *, verify_updates=True, stale=False):
    live = {}
    commands = []
    imported = []

    def adapter(*args):
        commands.append(args[0])
        if args[0] == "add-project-language":
            return ["LANGUAGE_ADDED|pl-PL"]
        if args[0] == "export-project-texts":
            if stale and "add-project-language" in commands:
                live["main-1"] = "Edited by another user"
            _export_file(Path(args[4]), live)
            return ["PROJECT_TEXTS_EXPORTED|" + args[4]]
        if args[0] == "import-project-texts":
            workbook = load_workbook(args[2])
            for row in list(workbook["User Texts"].values)[1:]:
                imported.append((row[2], row[5]))
                if verify_updates:
                    live[row[2]] = row[5]
            workbook.close()
            return [
                "PROJECT_TEXTS_IMPORTED|Success|C:\\TIA\\import.log",
                "PROJECT_SAVED|TestProject",
            ]
        raise AssertionError("Unexpected adapter call " + args[0])

    monkeypatch.setattr(
        review_module, "resolve_ollama_model",
        lambda requested: "test-local-model:latest",
    )
    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    monkeypatch.setattr(
        review_module, "_translate_texts",
        lambda texts, target, model, progress=None: {
            source: {"212BP2 Za zaworem": "212BP2 Za zaworem",
                     "Zgniatarka": "Zgniatarka"}.get(source, source)
            for source in texts
        },
    )
    review_module._REVIEWS.clear()
    return live, commands, imported


def test_prepare_does_not_modify_tia_and_shows_current_vs_proposed(monkeypatch):
    live, commands, imported = _setup(monkeypatch)
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["selected_rows"] == 2
    assert data["unique_texts"] == 2
    assert len(data["review_id"]) == 32
    assert data["examples"][0]["current_target"] == "Text"
    assert data["examples"][0]["proposed_target"] == "212BP2 Za zaworem"
    assert commands == ["export-project-texts"]
    assert imported == []
    assert live == {}


def test_import_only_approved_hmi_and_verify_export(monkeypatch):
    live, commands, imported = _setup(monkeypatch)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    ).json()
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "verified"
    assert response.json()["imported_rows"] == 2
    assert response.json()["verified_rows"] == 2
    assert imported == [
        ("main-1", "212BP2 Za zaworem"),
        ("main-2", "Zgniatarka"),
    ]
    assert "other-1" not in live
    assert commands == [
        "export-project-texts",   # prepare
        "add-project-language",   # only after approval
        "export-project-texts",   # re-check project just before import
        "import-project-texts",
        "export-project-texts",   # verify actual TIA values
    ]
    repeat = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert repeat.status_code == 404


def test_import_rejects_modified_tia_text_after_review(monkeypatch):
    _live, commands, imported = _setup(monkeypatch, stale=True)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    ).json()
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert response.status_code == 409, response.text
    assert imported == []
    assert "import-project-texts" not in commands


def test_import_is_not_reported_success_without_reexport_match(monkeypatch):
    _live, _commands, imported = _setup(monkeypatch, verify_updates=False)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    ).json()
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert len(imported) == 2
    assert response.status_code == 502
    assert "verification failed" in response.json()["detail"]
    assert "import.log" in response.json()["detail"]


def test_prepare_rejects_placeholder_for_nonplaceholder_source(monkeypatch):
    _setup(monkeypatch)
    monkeypatch.setattr(
        review_module, "_translate_texts",
        lambda texts, target, model, progress=None: {source: "Text" for source in texts},
    )
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert response.status_code == 502
    assert "placeholder" in response.json()["detail"]


def test_pilot_preserves_unreviewed_sibling_texts_in_same_hmi(monkeypatch):
    _live, _commands, imported = _setup(monkeypatch)
    pilot_request = {**REQUEST, "max_rows": 1}
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations",
        json=pilot_request,
    ).json()
    assert prepared["selected_rows"] == 1
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["verified_rows"] == 1
    # Parent-object/sibling rows remain in the XLSX with unchanged target.
    assert imported == [
        ("main-1", "212BP2 Za zaworem"),
        ("main-2", "Text"),
    ]


def test_translation_model_discovery_endpoint(monkeypatch):
    monkeypatch.setattr(
        review_module, "list_ollama_models",
        lambda: ["gemma3:4b", "qwen2.5:7b"],
    )
    monkeypatch.setattr(
        review_module, "resolve_ollama_model",
        lambda available: None if len(available) > 1 else available[0],
    )
    monkeypatch.setattr(review_module, "translation_batch_size", lambda: 1)
    response = client.get("/api/v1/tia/hmis/translation-models")
    assert response.status_code == 200
    assert response.json() == {
        "provider": "ollama",
        "models": ["gemma3:4b", "qwen2.5:7b"],
        "selected_model": None,
        "request_batch_size": 1,
    }


def test_preparation_stops_if_ollama_offline_before_export(monkeypatch):
    from fastapi import HTTPException

    review_module._REVIEWS.clear()

    def unavailable(_requested):
        raise HTTPException(503, "Cannot connect to local Ollama")

    def forbidden_adapter(*args):
        raise AssertionError("Preparing must not modify TIA when Ollama is offline")

    monkeypatch.setattr(review_module, "resolve_ollama_model", unavailable)
    monkeypatch.setattr(tia_openness, "_run_adapter", forbidden_adapter)
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert response.status_code == 503
    assert "Ollama" in response.json()["detail"]


def test_preparation_progress_endpoint(monkeypatch):
    from uuid import uuid4

    _setup(monkeypatch)
    progress_id = uuid4().hex
    request = {**REQUEST, "progress_id": progress_id}

    def translated(texts, target, model, progress=None):
        assert progress is not None
        progress(len(texts), len(texts))
        return {text: text for text in texts}

    monkeypatch.setattr(review_module, "_translate_texts", translated)
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=request,
    )
    assert response.status_code == 200, response.text

    status = client.get(
        "/api/v1/tia/hmis/translation-progress/" + progress_id,
    )
    assert status.status_code == 200
    data = status.json()
    assert data["stage"] == "ready"
    assert data["total"] == 2
    assert data["completed"] == 2
    assert data["percent"] == 100.0
    assert data["elapsed_seconds"] >= 0


def test_preparation_progress_rejects_bad_identifier():
    response = client.get(
        "/api/v1/tia/hmis/translation-progress/not-valid",
    )
    assert response.status_code == 400


def test_manual_prepare_rejects_unbounded_requests_before_export(monkeypatch):
    def forbidden(*arguments):
        raise AssertionError("Invalid manual batch must never export or change TIA")

    monkeypatch.setattr(tia_openness, "_run_adapter", forbidden)
    for amount in (None, 51):
        response = client.post(
            "/api/v1/tia/hmis/prepare-translations",
            json={**REQUEST, "max_rows": amount},
        )
        assert response.status_code == 422, response.text
        assert "max_rows from 1 to 50" in response.json()["detail"]


def test_manual_prepare_uses_requested_batch_size(monkeypatch):
    _, commands, _ = _setup(monkeypatch)
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations",
        json={**REQUEST, "max_rows": 1},
    )
    assert response.status_code == 200, response.text
    assert response.json()["selected_rows"] == 1
    assert response.json()["unique_texts"] == 1
    assert commands == ["export-project-texts"]


def test_prepare_skips_populated_targets_to_preserve_existing_polish(monkeypatch):
    live, commands, imported = _setup(monkeypatch)
    live["main-2"] = "Zgniatarka"  # already correct in TIA
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations",
        json={**REQUEST, "max_rows": 10},
    )
    assert response.status_code == 200, response.text
    prepared = response.json()
    assert prepared["unique_texts"] == 1
    assert prepared["examples"][0]["source_text"] == "212BP2 Za zaworem"
    saved = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared["review_id"]},
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["verified_rows"] == 1
    assert live["main-2"] == "Zgniatarka"
    assert "import-project-texts" in commands


def test_source_copy_does_not_call_ollama_and_preserves_label(monkeypatch):
    live, commands, imported = _setup(monkeypatch)
    def forbidden_ollama(*args, **kwargs):
        raise AssertionError("Exact source-copy mode should not call Ollama")
    monkeypatch.setattr(review_module, "_translate_texts", forbidden_ollama)
    monkeypatch.setattr(review_module, "resolve_ollama_model", forbidden_ollama)
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations",
        json={**REQUEST, "translation_strategy": "copy", "max_rows": 1},
    )
    assert response.status_code == 200, response.text
    assert response.json()["examples"][0]["proposed_target"] == "212BP2 Za zaworem"
    assert commands == ["export-project-texts"]


def test_import_info_state_is_verified_after_reexport(monkeypatch):
    """TIA Openness may report 'Info' for a completed project-text import."""
    live, commands, imported = _setup(monkeypatch)
    original_adapter = tia_openness._run_adapter

    def adapter(*args):
        result = original_adapter(*args)
        if args[0] == "import-project-texts":
            return [
                line.replace("PROJECT_TEXTS_IMPORTED|Success|",
                             "PROJECT_TEXTS_IMPORTED|Info|", 1)
                for line in result
            ]
        return result

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert prepared.status_code == 200, prepared.text
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared.json()["review_id"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "verified"
    assert response.json()["import_state"] == "Info"
    assert response.json()["verified_rows"] == 2
    assert commands.count("export-project-texts") == 3
    assert commands.count("import-project-texts") == 1


def test_info_state_cannot_bypass_failed_reexport_verification(monkeypatch):
    _live, commands, imported = _setup(monkeypatch, verify_updates=False)
    original_adapter = tia_openness._run_adapter

    def adapter(*args):
        result = original_adapter(*args)
        if args[0] == "import-project-texts":
            return [
                line.replace("PROJECT_TEXTS_IMPORTED|Success|",
                             "PROJECT_TEXTS_IMPORTED|Info|", 1)
                for line in result
            ]
        return result

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert prepared.status_code == 200, prepared.text
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared.json()["review_id"]},
    )
    assert response.status_code == 502
    assert "re-export verification failed" in response.json()["detail"]
    assert commands.count("export-project-texts") == 3


def test_import_warning_still_stops_before_reporting_success(monkeypatch):
    _live, commands, imported = _setup(monkeypatch)
    original_adapter = tia_openness._run_adapter

    def adapter(*args):
        result = original_adapter(*args)
        if args[0] == "import-project-texts":
            return [
                line.replace("PROJECT_TEXTS_IMPORTED|Success|",
                             "PROJECT_TEXTS_IMPORTED|Warning|", 1)
                for line in result
            ]
        return result

    monkeypatch.setattr(tia_openness, "_run_adapter", adapter)
    prepared = client.post(
        "/api/v1/tia/hmis/prepare-translations", json=REQUEST,
    )
    assert prepared.status_code == 200, prepared.text
    response = client.post(
        "/api/v1/tia/hmis/import-reviewed-translations",
        json={"review_id": prepared.json()["review_id"]},
    )
    assert response.status_code == 502
    assert "Warning" in response.json()["detail"]
    assert commands.count("export-project-texts") == 2



def test_manual_prepare_normalizes_literal_backslash_n_before_review(
    monkeypatch,
):
    _live, commands, imported = _setup(monkeypatch)
    monkeypatch.setattr(
        review_module, "_translate_texts",
        lambda texts, target, model, progress=None: {
            text: r"212BP2 Za\nzaworem" if text == "212BP2 Za zaworem"
            else text for text in texts
        },
    )
    response = client.post(
        "/api/v1/tia/hmis/prepare-translations",
        json=REQUEST,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["examples"][0]["proposed_target"] == "212BP2 Za zaworem"
    assert commands == ["export-project-texts"]
    assert not imported
