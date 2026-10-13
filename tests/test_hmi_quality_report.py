"""Offline XLSX export contract for HMI translation quality reports."""
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from engine.main import app
from engine.tia import hmi_batches, hmi_quality_report as report
from engine.tia.hmi_review import (
    ExistingHmiQaAuditResponse,
    ExistingHmiQaFinding,
)


client = TestClient(app)
JOB_ID = "a" * 32


def fake_job():
    return {
        "job_id": JOB_ID,
        "request": {
            "hmi_name": "HMI_RT_1",
            "view_path_root": "HMI_00_Main",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
        },
        "model": "qwen2.5:3b",
        "batch_size": 2,
        "candidates": [
            "Naciągnij linkę",
            "=SCADA unsafe label",
            "Awaria PLC",
            "Od strony karuzeli",
            "Pressure low",
        ],
        "next_offset": 4,
        "completed_rows": 7,
        "skipped": [
            {"offset": 1, "source": "=SCADA unsafe label",
             "proposal": "damaged (", "reason": "changed parentheses"},
            {"offset": 2, "source": "Awaria PLC",
             "proposal": "unknown", "reason": "old AI proposal invalid"},
        ],
        "verified_details": {
            "0": {
                "proposal": "Naciągnij linkę",
                "auto_repaired": True,
                "verified_rows": 2,
                "batch_index": 1,
            },
            "3": {
                "proposal": "Od strony karuzeli",
                "auto_repaired": False,
                "verified_rows": 2,
                "batch_index": 2,
            },
        },
        "recovered": {
            "2": {"proposal": "Awaria PLC", "verified_rows": 3,
                  "mode": "manual-reviewed"},
        },
        "status": "paused",
        "started_at": 1,
        "updated_at": 2,
    }


def _open(response):
    assert response.status_code == 200, response.text
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    return load_workbook(BytesIO(response.content), data_only=False)


def test_batch_quality_report_has_real_per_text_verified_outcomes(
    monkeypatch, tmp_path,
):
    monkeypatch.setattr(hmi_batches, "JOB_DIR", tmp_path)
    hmi_batches._save_job(fake_job())
    response = client.get(
        f"/api/v1/tia/hmis/batch-translations/{JOB_ID}/quality-report.xlsx"
    )
    workbook = _open(response)
    try:
        assert workbook.sheetnames == ["Summary", "Translations"]
        assert workbook["Summary"]["B2"].value == JOB_ID
        rows = list(workbook["Translations"].values)
        assert len(rows) == 6  # header plus every source, including unprocessed
        columns = {name: i for i, name in enumerate(rows[0])}
        status = columns["Status"]
        proposed = columns["Approved or rejected proposal"]
        original = columns["Source text"]
        assert rows[1][status] == "Auto-repaired — verified in TIA"
        assert rows[1][proposed] == "Naciągnij linkę"
        assert rows[2][status] == "Skipped — not imported"
        assert rows[2][original] == "'=SCADA unsafe label"
        assert rows[2][proposed] == "damaged ("
        assert rows[3][status] == "Recovered — verified in TIA"
        assert rows[3][proposed] == "Awaria PLC"
        assert rows[4][status] == "Verified in TIA"
        assert rows[4][proposed] == "Od strony karuzeli"
        assert rows[5][status] == "Not processed"
        # Untrusted text cannot become an active Excel formula.
        assert workbook["Translations"]["C3"].data_type == "s"
    finally:
        workbook.close()


def test_legacy_checkpoint_reports_unknown_proposal_without_invention(
    monkeypatch, tmp_path,
):
    monkeypatch.setattr(hmi_batches, "JOB_DIR", tmp_path)
    old = fake_job()
    del old["verified_details"]
    del old["recovered"]
    old["next_offset"] = 1
    old["skipped"] = []
    old["status"] = "error"
    hmi_batches._save_job(old)
    workbook = _open(client.get(
        f"/api/v1/tia/hmis/batch-translations/{JOB_ID}/quality-report.xlsx"
    ))
    try:
        assert workbook["Translations"]["B2"].value == (
            "Verified — legacy checkpoint"
        )
        assert workbook["Translations"]["D2"].value is None
        assert workbook["Translations"]["B3"].value == (
            "Not processed — previous import state uncertain"
        )
    finally:
        workbook.close()


def test_existing_target_quality_xlsx_contains_full_audit_findings(monkeypatch):
    audit_calls = []

    def fake_audit(request, *, limit):
        audit_calls.append((request.hmi_name, limit))
        return ExistingHmiQaAuditResponse(
            hmi_name=request.hmi_name,
            scanned_rows=500,
            filled_target_rows=380,
            suspicious_rows=2,
            shown_rows=2,
            limit=limit,
            findings=[
                ExistingHmiQaFinding(
                    sheet="User Texts", row=23,
                    view_path="Project\\HMI_00_Main\\Screens\\Main",
                    source_text="(Hydrauliczna)", target_text="(cykliski)",
                    reasons=["Unexpected malformed word: 'cykliski'."],
                ),
                ExistingHmiQaFinding(
                    sheet="User Texts", row=24,
                    view_path="Project\\HMI_00_Main\\Text lists\\Auto Status",
                    source_text="Naciągnij linkę", target_text="Naciągnij linię",
                    reasons=["Technical meaning changed: linka → linia"],
                ),
            ],
        )

    monkeypatch.setattr(
        report, "_audit_existing_hmi_translations", fake_audit,
    )
    response = client.post(
        "/api/v1/tia/hmis/audit-existing-translations/quality-report.xlsx",
        json={
            "process_id": 1337,
            "device_name": "HMI_RT_1",
            "hmi_name": "HMI_RT_1",
            "view_path_root": "HMI_00_Main",
            "source_culture": "en-US",
            "target_culture": "pl-PL",
        },
    )
    workbook = _open(response)
    try:
        assert workbook.sheetnames == ["Summary", "Existing QA"]
        assert audit_calls == [("HMI_RT_1", 1000000)]
        entries = list(workbook["Existing QA"].values)
        assert len(entries) == 3
        assert entries[1][4] == "(cykliski)"
        assert entries[2][4] == "Naciągnij linię"
        assert "Text lists" in entries[2][2]
        assert workbook["Existing QA"].freeze_panes == "A2"
    finally:
        workbook.close()


def test_invalid_batch_id_cannot_escape_checkpoint_directory(monkeypatch, tmp_path):
    monkeypatch.setattr(hmi_batches, "JOB_DIR", tmp_path)
    response = client.get(
        "/api/v1/tia/hmis/batch-translations/invalid/quality-report.xlsx"
    )
    assert response.status_code == 422
